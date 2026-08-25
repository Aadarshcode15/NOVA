# actions/file_analyzer.py
import base64
import io
import threading
from pathlib import Path
from core.logger import log
from core.voice import speak
from core.engine import raw_query, query_vision_groq


# ── Session state ──────────────────────────────────────────────
_active: dict = {
    "path":    None,
    "type":    None,
    "name":    "",
    "content": "",
    "lang":    "",
}

def get_active() -> dict:  return _active
def has_active_file() -> bool: return _active["path"] is not None


# ── File type detection ────────────────────────────────────────
_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
_CODE_EXT  = {".py", ".js", ".ts", ".html", ".css", ".java",
              ".cpp", ".c", ".go", ".rb", ".sql", ".sh", ".kt"}
_TEXT_EXT  = {".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml"}
_PDF_EXT   = {".pdf"}
_WORD_EXT  = {".docx", ".doc"}
_EXCEL_EXT = {".xlsx", ".xls", ".xlsm"}

_LANG_NAMES = {
    ".py": "Python", ".js": "JavaScript", ".ts": "TypeScript",
    ".html": "HTML",  ".css": "CSS",       ".java": "Java",
    ".cpp": "C++",    ".c": "C",           ".go": "Go",
    ".rb": "Ruby",    ".sql": "SQL",       ".sh": "Bash",
    ".kt": "Kotlin",
}

def _detect_type(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in _IMAGE_EXT:  return "image"
    if ext in _PDF_EXT:    return "pdf"
    if ext in _WORD_EXT:   return "word"
    if ext in _EXCEL_EXT:  return "excel"
    if ext in _CODE_EXT:   return "code"
    if ext in _TEXT_EXT:   return "text"
    return "unknown"


# ── Tesseract path (Windows) ───────────────────────────────────
def _configure_tesseract() -> bool:
    """Configure pytesseract to find the Tesseract binary on Windows."""
    try:
        import pytesseract
        import subprocess
        # Try running tesseract to see if it's in PATH
        subprocess.run(
            ["tesseract", "--version"],
            capture_output=True, timeout=3
        )
        return True
    except FileNotFoundError:
        # Common Windows install location
        import os
        common_path = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        if os.path.exists(common_path):
            import pytesseract
            pytesseract.pytesseract.tesseract_cmd = common_path
            return True
        log.warning(
            "[OCR] Tesseract not found. Install from: "
            "https://github.com/UB-Mannheim/tesseract/wiki"
        )
        return False
    except Exception:
        return False


# ── PDF extraction ─────────────────────────────────────────────

def _read_pdf_text(path: Path) -> tuple:
    """
    Extract text from PDF. Uses pdfplumber (better than pypdf) first.
    Falls back to pypdf if pdfplumber fails.
    Returns (text, page_count).
    """
    # ── Method 1: pdfplumber (handles most modern PDFs) ──
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            pages = len(pdf.pages)
            parts = []
            for page in pdf.pages[:20]:
                t = page.extract_text()
                if t:
                    parts.append(t.strip())
            text = "\n\n".join(parts).strip()
            if text:
                return text, pages
    except Exception as e:
        log.debug(f"[FileAnalyzer] pdfplumber failed: {e}")

    # ── Method 2: pypdf fallback ──
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages  = len(reader.pages)
        parts  = []
        for page in reader.pages[:20]:
            t = page.extract_text()
            if t:
                parts.append(t.strip())
        text = "\n".join(parts).strip()
        if text:
            return text, pages
    except Exception as e:
        log.debug(f"[FileAnalyzer] pypdf failed: {e}")

    # If both fail, return empty — OCR will be tried next
    return "", 0


def _ocr_pdf(path: Path) -> str:
    """
    OCR fallback for scanned / image-based PDFs.
    Uses PyMuPDF to render pages as images, then Tesseract to read them.
    """
    if not _configure_tesseract():
        return ""
    try:
        import fitz          # PyMuPDF
        import pytesseract
        from PIL import Image

        doc   = fitz.open(str(path))
        parts = []
        for page_num in range(min(10, len(doc))):
            page   = doc[page_num]
            # 2x zoom gives better OCR accuracy
            matrix = fitz.Matrix(2.0, 2.0)
            pix    = page.get_pixmap(matrix=matrix, colorspace=fitz.csGRAY)
            img    = Image.frombytes("L", [pix.width, pix.height], pix.samples)
            text   = pytesseract.image_to_string(img, config="--psm 6")
            if text.strip():
                parts.append(text.strip())
        doc.close()
        return "\n\n".join(parts)
    except ImportError:
        log.warning("[OCR] PyMuPDF or pytesseract not installed.")
        return ""
    except Exception as e:
        log.error(f"[OCR] PDF OCR error: {e}")
        return ""


def _read_pdf(path: Path) -> tuple:
    """Main PDF reader: text extraction → OCR fallback."""
    text, pages = _read_pdf_text(path)
    if text:
        return text, pages, False   # (text, pages, used_ocr)

    # Text extraction failed → this is likely a scanned PDF
    log.info(f"[FileAnalyzer] No text layer found in {path.name} — trying OCR")
    ocr_text = _ocr_pdf(path)
    if ocr_text:
        # Get page count even if text was OCR'd
        try:
            import fitz
            doc   = fitz.open(str(path))
            pages = len(doc)
            doc.close()
        except Exception:
            pages = 1
        return ocr_text, pages, True
    return "", 0, False


# ── Word document extraction ───────────────────────────────────

def _read_word(path: Path) -> tuple:
    """Extract text from .docx Word documents including tables."""
    try:
        from docx import Document
        doc   = Document(str(path))
        parts = []

        # Paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())

        # Tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    parts.append(" | ".join(row_cells))

        text = "\n".join(parts)
        return text.strip(), len(doc.paragraphs)

    except ImportError:
        log.warning("[FileAnalyzer] python-docx not installed — pip install python-docx")
        return "", 0
    except Exception as e:
        log.error(f"[FileAnalyzer] Word read error: {e}")
        return "", 0


# ── Excel extraction ───────────────────────────────────────────

def _read_excel(path: Path) -> tuple:
    """Extract data from .xlsx Excel files as readable text."""
    try:
        import openpyxl
        wb     = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        parts  = []
        sheets = 0
        rows_total = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"[Sheet: {sheet_name}]")
            for row in ws.iter_rows(values_only=True):
                row_text = "  |  ".join(
                    str(v) for v in row if v is not None and str(v).strip()
                )
                if row_text.strip():
                    parts.append(row_text)
                    rows_total += 1
            sheets += 1
            if rows_total > 500:   # cap to avoid token overflow
                parts.append("... (truncated — too many rows)")
                break

        wb.close()
        return "\n".join(parts), sheets

    except ImportError:
        log.warning("[FileAnalyzer] openpyxl not installed — pip install openpyxl")
        return "", 0
    except Exception as e:
        log.error(f"[FileAnalyzer] Excel read error: {e}")
        return "", 0


# ── Code / Text extraction ─────────────────────────────────────

def _read_code(path: Path) -> tuple:
    content = path.read_text(encoding="utf-8", errors="ignore")
    lines   = len(content.splitlines())
    lang    = _LANG_NAMES.get(path.suffix.lower(), path.suffix[1:].upper())
    return content, lines, lang

def _read_text(path: Path) -> tuple:
    content = path.read_text(encoding="utf-8", errors="ignore")
    words   = len(content.split())
    return content, words


# ── OCR for images ─────────────────────────────────────────────

def _ocr_image(path: Path) -> str:
    """Run OCR on an image file to extract text."""
    if not _configure_tesseract():
        return ""
    try:
        import pytesseract
        from PIL import Image
        img  = Image.open(str(path))
        text = pytesseract.image_to_string(img, config="--psm 6")
        return text.strip()
    except Exception as e:
        log.error(f"[OCR] Image OCR error: {e}")
        return ""


# ── Main: on-upload analysis ───────────────────────────────────

def load_and_announce(filepath: str) -> None:
    """
    Called in background thread when user uploads a file.
    Detects type, extracts content (with OCR fallback), and speaks result.
    """
    global _active
    path     = Path(filepath)
    filetype = _detect_type(path)

    _active = {"path": path, "type": filetype, "name": path.name,
               "content": "", "lang": ""}

    log.info(f"[FileAnalyzer] Loaded: {path.name} ({filetype})")

    # ── IMAGE ──────────────────────────────────────────────────
    if filetype == "image":
        speak(f"Let me take a look at {path.name}.")
        try:
            with open(path, "rb") as f:
                img_b64 = base64.b64encode(f.read()).decode()

            description = query_vision_groq(
                img_b64,
                "Describe what you see in this image in 2 to 3 clear spoken sentences. "
                "Be specific. If there is text visible, mention what it says. "
                "No markdown, no bullet points."
            )
            _active["content"] = description
            speak(description)

            # Check if image has extractable text
            ocr_text = _ocr_image(path)
            if ocr_text and len(ocr_text) > 20:
                _active["content"] = ocr_text
                speak(
                    f"I also extracted {len(ocr_text.split())} words of text from "
                    f"this image. Say 'read the text' to hear it, or "
                    f"'summarize it' for a summary."
                )
            else:
                speak("Say 'describe it in more detail' if you want more, "
                      "or ask me anything about this image.")

        except Exception as e:
            log.error(f"[FileAnalyzer] Image analysis error: {e}")
            speak(f"I loaded {path.name} but had trouble analyzing it.")
        return

    # ── PDF ────────────────────────────────────────────────────
    if filetype == "pdf":
        speak(f"Reading {path.name}. One moment.")
        text, pages, used_ocr = _read_pdf(path)
        _active["content"] = text
        if not text:
            speak(
                f"I loaded {path.name} but couldn't extract any text. "
                f"The file may be corrupted or encrypted."
            )
            return
        ocr_note = " using OCR" if used_ocr else ""
        preview  = text[:120].replace("\n", " ").strip()
        speak(
            f"Loaded {path.name}{ocr_note} — {pages} page{'s' if pages != 1 else ''}, "
            f"{len(text.split())} words extracted. "
            f"It begins: {preview}... "
            f"Say 'summarize it' for key points."
        )
        return

    # ── WORD ───────────────────────────────────────────────────
    if filetype == "word":
        speak(f"Reading {path.name}.")
        text, para_count = _read_word(path)
        _active["content"] = text
        if not text:
            speak(f"I couldn't extract text from {path.name}. "
                  "Make sure python-docx is installed.")
            return
        preview = text[:120].replace("\n", " ").strip()
        speak(
            f"Loaded {path.name} — {para_count} paragraphs, "
            f"{len(text.split())} words. "
            f"It starts with: {preview}... "
            f"Say 'summarize it' for a summary."
        )
        return

    # ── EXCEL ──────────────────────────────────────────────────
    if filetype == "excel":
        speak(f"Reading {path.name}.")
        text, sheet_count = _read_excel(path)
        _active["content"] = text
        if not text:
            speak(f"I couldn't read {path.name}. "
                  "Make sure openpyxl is installed.")
            return
        rows = text.count("\n")
        speak(
            f"Loaded {path.name} — {sheet_count} sheet{'s' if sheet_count != 1 else ''}, "
            f"approximately {rows} rows of data. "
            f"Say 'summarize it' for an overview of the data."
        )
        return

    # ── CODE ───────────────────────────────────────────────────
    if filetype == "code":
        content, lines, lang = _read_code(path)
        _active["content"]   = content
        _active["lang"]      = lang
        speak(
            f"Loaded {path.name} — {lang} file, {lines} lines. "
            f"Say 'explain it' to understand what it does, "
            f"or 'fix it' to debug any issues."
        )
        return

    # ── TEXT ───────────────────────────────────────────────────
    if filetype == "text":
        content, words = _read_text(path)
        _active["content"] = content
        speak(
            f"Loaded {path.name} — {words} words. "
            f"Say 'summarize it' for an overview."
        )
        return

    speak(
        f"I loaded {path.name}. I can work with images, PDFs, "
        f"Word documents, Excel files, code files, and text files."
    )


# ── File command handler (called from router) ──────────────────

_SUMMARIZE = (
    "summarize it", "summarize the file", "summarize this",
    "give me a summary", "what does it say", "what's in it",
    "sum it up", "give me the gist", "what is this about",
    "what's this file about", "overview",
)
_EXPLAIN = (
    "explain it", "explain the code", "what does this code do",
    "what does it do", "break it down", "what is this code",
    "walk me through it",
)
_DETAIL = (
    "more detail", "tell me more", "describe it in more detail",
    "go deeper", "read the text", "read it",
)
_FIX = (
    "fix it", "debug it", "fix the code", "debug the code",
    "find the bugs",
)


def handle_file_command(command: str) -> bool:
    if not has_active_file():
        return False

    c       = command.lower().strip()
    name    = _active["name"]
    typ     = _active["type"]
    content = _active["content"]

    # ── Summarize ──────────────────────────────────────────────
    if any(t in c for t in _SUMMARIZE):
        if not content:
            speak(f"I couldn't extract text from {name}. "
                  "It may be a scanned or protected file.")
            return True
        speak(f"Summarizing {name}.")
        summary = raw_query(
            f"Summarize this document in 3 concise spoken sentences. "
            f"Focus on the main topic and most important points. "
            f"No markdown, no bullet points:\n\n{content[:4000]}",
            system=(
                "You are a document summarizer for a voice assistant. "
                "Return only a brief spoken summary. No formatting."
            )
        )
        speak(summary or "Could not generate summary.")
        return True

    # ── Explain code ───────────────────────────────────────────
    if any(t in c for t in _EXPLAIN):
        if typ != "code":
            speak(f"{name} is not a code file. Say 'summarize it' instead.")
            return True
        lang = _active.get("lang", "code")
        speak(f"Explaining the {lang} code.")
        explanation = raw_query(
            f"Explain what this {lang} code does in 3 spoken sentences. "
            f"Focus on WHAT it does and WHY, not the syntax. "
            f"No markdown:\n\n{content[:3000]}",
            system="You are a code explainer for a voice assistant. Be clear and concise."
        )
        speak(explanation or "Could not explain the code.")
        return True

    # ── Read text / more detail ────────────────────────────────
    if any(t in c for t in _DETAIL):
        if typ == "image" and content:
            try:
                with open(_active["path"], "rb") as f:
                    img_b64 = base64.b64encode(f.read()).decode()
                detail = query_vision_groq(
                    img_b64,
                    "Give a detailed description of this image. Include any visible text, "
                    "colours, layout, numbers, and all important details. 3-4 sentences, no markdown."
                )
                speak(detail)
            except Exception as e:
                log.error(f"[FileAnalyzer] Detail vision error: {e}")
                speak("Couldn't get more detail right now.")
        elif content:
            # Read first 400 chars aloud
            snippet = content[:400].replace("\n", " ").strip()
            speak(snippet)
        else:
            speak(f"No content available for {name}.")
        return True

    # ── Fix code ───────────────────────────────────────────────
    if any(t in c for t in _FIX):
        if typ != "code":
            speak(f"{name} is not a code file.")
            return True
        from core.engine import raw_query as rq
        fixed = rq(
            f"Fix ALL bugs in this code. Return ONLY fixed code:\n\n{content}"
        )
        if fixed:
            backup = _active["path"].parent / (_active["path"].stem + "_backup" + _active["path"].suffix)
            backup.write_text(content, encoding="utf-8")
            _active["path"].write_text(fixed, encoding="utf-8")
            speak(f"Fixed and saved. Original backed up as {backup.name}.")
        else:
            speak("Could not fix the code right now.")
        return True

    # ── Pass-through: any question about the file ──────────────
    if content and any(k in c for k in ("this file", "the file", "this document",
                                         "this code", "the code", "it", "this")):
        from core.engine import query
        response = query(
            f"File: '{name}'\nContent:\n{content[:3000]}\n\nQuestion: {command}",
            skip_history=True
        )
        speak(response)
        return True

    return False